"""Ungeprüfte Spitzenreiter außerhalb der USA gegen die Kursarchive ihrer Börsen (ab 1.19.0, Prüfskript; App unverändert).

Simons Auftrag vom 25.10.2026, 14:23 Uhr, Punkt 1: „Prüfe die bisher ungeprüften Spitzenreiter, insbesondere Warschau und TSX
Venture.“ Bis dahin verglich der Job „welt“ Aktien außerhalb der US-Auswahl nur mit CNBC; für Warschau fand CNBC keine der sechs
Aktien mit Platz unter den ersten 10 (16 Plätze 1), für TSX Venture war die Zuordnung der Aktie mit 22 Plätzen 1 nicht
abgesichert (Kursniveau Stand/CNBC 1,059). Hier alle Aktien außerhalb der US-Auswahl mit einem Platz unter den ersten 10 an
diesen Börsen gegen die Quellen der Börsenbetreiber (frei, ohne Anmeldung):

* **Warschau (XWAR):** Notierungsarchiv der GPW je Handelstag, volle Tabelle – alle Aktien des Hauptmarkts mit ISIN, Währung,
  Schlusskurs und Umsatz (``GPW_URL``). Die ISIN macht die Zuordnung eindeutig.
* **Toronto und TSX Venture (XTSE, XTSX):** Kurshistorie bei TMX Money (die TMX Group betreibt beide Börsen), Schnittstelle der
  Website (``TMX_URL``) – Name und Börse aus der Kursmeldung, Aktie und Gattung zusätzlich über OpenFIGI. Die Währung folgt der
  Kürzelregel der TMX (Kurse in CAD, Kürzel mit „.U“ in USD) und wird zusätzlich über das Kursniveau geprüft.

Je Aktie dieselben strengen Prüfungen wie im Job „welt“ (``kurse_welt.check_world``) und dieselbe Einordnung jedes Tages
(``kurse_welt.classify_days``); für abgesicherte Aktien mit abweichenden Tagen alle Ranglisten mit den Kursen der Börse neu
gerechnet (``kurse_ursachen.ranking_effect``) – Platz 1 und Plätze unter den ersten 10 vorher/nachher. Öffentlich nur Zählwerte,
Börsencodes, Buchstaben und Fundstellen (Zeitraum/Gruppe/Gewichtung/Platz) – keine Kurse, Namen, Kürzel oder ISINs. Der Stand
wird nicht geändert; der Job wird rot, solange ein Spitzenreiter ungeprüft oder ein Kurs nicht bestätigt ist.

Dazu (``fx_probe``): Erreichbarkeit frei zugänglicher Quellen für Taiwan-Dollar und Chile-Peso (die EZB führt beide nicht) –
Fed H.10 über FRED, BIS, mindicador.cl (Daten der Banco Central de Chile) – nur Status, letzter Tag und Zahl der Tage.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import os
import re
import sys
import tempfile
import time
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parents[1] / "backend"))

import kurse_absicherung as absicherung  # noqa: E402
import kurse_ursachen as ursachen  # noqa: E402
import kurse_welt as welt  # noqa: E402
import nachrechnen  # noqa: E402

EPOCH = dt.date(1970, 1, 1)
# Volle Tabelle des Archivs (mit „Kod ISIN“ und Umsatz; die kurze unter /archiwum-notowan hat keine ISIN – Lauf 40201584212)
GPW_URL = "https://www.gpw.pl/archiwum-notowan-full?type=10&instrument=&date={date}"
GPW_HEADERS = {"Accept": "text/html,application/xhtml+xml", "Accept-Language": "pl,en;q=0.8", "Referer": "https://www.gpw.pl/archiwum-notowan-full"}
# NewConnect (alternativer Markt der GPW): keine Archivtabelle, aber eine Tagesstatistik je Handelstag als Datei (seit Oktober
# 2026 .xlsx, davor .xls) – Lauf 40202290066: zwei Warschauer Aktien nicht im Archiv des Hauptmarkts
NC_URL = "https://newconnect.pl/pub/NEWCONNECT/statystyki/statystyki_dzienne/NC_{ymd}.{ext}"
TMX_URL = "https://app-money.tmx.com/graphql"
TMX_HEADERS = {"Accept": "*/*", "Content-Type": "application/json", "locale": "en", "Origin": "https://money.tmx.com", "Referer": "https://money.tmx.com/"}
HOSTS = {"www.gpw.pl": 0.5, "newconnect.pl": 0.5, "app-money.tmx.com": 0.5, "fred.stlouisfed.org": 0.5, "stats.bis.org": 0.5, "mindicador.cl": 0.5,
         "query2.finance.yahoo.com": 0.3, "api.openfigi.com": 2.5}
SOURCES = {"XWAR": "gpw", "XTSE": "tmx", "XTSX": "tmx"}
SOURCE_TEXT = {"gpw": "Notierungsarchiv der GPW", "tmx": "TMX Money", "newconnect": "Tagesstatistik der NewConnect (GPW)"}
VENUE_NAME = {"gpw": "Warsaw Stock Exchange"}  # die GPW nennt im Archiv keine Börse – es ist ihr eigenes Archiv
ISIN_RE = re.compile(r"[A-Z]{2}[A-Z0-9]{9}\d")

TMX_QUOTE = (
    "query getQuoteBySymbol($symbol: String, $locale: String) {"
    " getQuoteBySymbol(symbol: $symbol, locale: $locale) { symbol name exchangeName exchangeCode } }"
)
TMX_HISTORY = (
    "query getCompanyPriceHistory($symbol: String!, $start: String, $end: String, $limit: Int, $adjusted: Boolean,"
    " $adjustmentType: String, $unadjusted: Boolean) { getCompanyPriceHistory(symbol: $symbol, start: $start, end: $end,"
    " limit: $limit, adjusted: $adjusted, adjustmentType: $adjustmentType, unadjusted: $unadjusted) { datetime closePrice volume } }"
)
# Ab dem Lauf 40201584212: getCompanyPriceHistory lieferte je Abfrage nur die letzten 25 Handelstage – deshalb zuerst die
# Tageskerzen der Kursgrafik (getTimeSeriesData), sonst die Kurshistorie in Abschnitten rückwärts bis zum Beginn des Zeitraums.
TMX_SERIES = (
    "query getTimeSeriesData($symbol: String!, $freq: String, $interval: Int, $start: String, $end: String, $startDateTime: Int,"
    " $endDateTime: Int) { getTimeSeriesData(symbol: $symbol, freq: $freq, interval: $interval, start: $start, end: $end,"
    " startDateTime: $startDateTime, endDateTime: $endDateTime) { dateTime close volume } }"
)
TMX_PAGES = 8


def iso(day: int) -> str:
    return (EPOCH + dt.timedelta(days=day)).isoformat()


class _Tables(HTMLParser):
    """Alle Tabellen einer Seite als Zeilen von Zelltexten (Leerraum zusammengefasst)."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self._row and self.tables:
                self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def pl_number(text: str | None) -> float | None:
    """Zahl in polnischer Schreibweise („1 234,56“, auch mit geschütztem Leerzeichen); „-“ oder leer → None."""
    s = (text or "").replace("\xa0", "").replace(" ", "").replace(",", ".").strip()
    if not s or s in ("-", "–", "—"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_gpw(html: str) -> tuple[str, dict[str, dict], list[str]]:
    """Notierungsarchiv eines Tages: (Status, ISIN → Schlusskurs, Umsatz, Währung, Name; Spaltenköpfe ohne Werte)."""
    parser = _Tables()
    parser.feed(html or "")
    for table in parser.tables:
        head_at = next((k for k, row in enumerate(table) if any("isin" in c.lower() for c in row)), None)
        if head_at is None:
            continue
        header = [c.lower() for c in table[head_at]]

        def col(*words: str) -> int | None:
            return next((k for k, c in enumerate(header) if any(w in c for w in words)), None)

        c_isin, c_close = col("isin"), col("zamkn", "closing", "close")
        c_cur, c_vol, c_name = col("walut", "currency"), col("wolumen", "volume"), col("nazwa", "name")
        if c_isin is None or c_close is None:
            return "spalten_fehlen", {}, table[head_at]
        out: dict[str, dict] = {}
        for row in table[head_at + 1:]:
            if len(row) <= max(c_isin, c_close):
                continue
            isin = row[c_isin].strip().upper()
            close = pl_number(row[c_close])
            if not ISIN_RE.fullmatch(isin) or close is None or close <= 0:
                continue
            out[isin] = {
                "close": close,
                "volume": pl_number(row[c_vol]) if c_vol is not None and c_vol < len(row) else None,
                "currency": row[c_cur].strip() if c_cur is not None and c_cur < len(row) else None,
                "name": row[c_name].strip() if c_name is not None and c_name < len(row) else None,
            }
        return ("ok" if out else "ohne_zeilen"), out, table[head_at]
    heads = [" / ".join(table[0][:4]) for table in parser.tables if table][:3]
    return "ohne_tabelle", {}, [f"{len(parser.tables)} Tabellen ohne ISIN-Spalte"] + heads


def fetch_gpw(http, days: list[int], pause=time.sleep) -> tuple[Counter, dict[int, dict[str, dict]], list[str]]:
    """Archiv je Handelstag; ein Tag ohne Antwort wird nach ``welt.RETRY_PAUSE_S`` einmal wiederholt. Liefern die ersten drei
    Tage nur Unlesbares oder Fehler, bricht der Abruf ab (Format oder Zugang geändert – gezählt als „abgebrochen“)."""
    status: Counter = Counter()
    by_day: dict[int, dict[str, dict]] = {}
    header: list[str] = []
    for n, day in enumerate(days):
        if n >= 3 and not by_day:
            status["abgebrochen_nach_3_tagen_ohne_daten"] += len(days) - n
            break
        url = GPW_URL.format(date=(EPOCH + dt.timedelta(days=day)).strftime("%d-%m-%Y"))
        code, content = welt.zusatz.get_raw(http, url, GPW_HEADERS)
        if code != "ok":
            pause(welt.RETRY_PAUSE_S)
            code, content = welt.zusatz.get_raw(http, url, GPW_HEADERS)
        if code != "ok":
            status[code] += 1
            continue
        result, rows, head = parse_gpw((content or b"").decode("utf-8", errors="replace"))
        header = header or head
        status[result] += 1
        if result == "ok":
            by_day[day] = rows
    return status, by_day, header


def parse_rows(table: list[list[str]]) -> tuple[str, dict[str, dict], list[str]]:
    """Zeilen einer Tabelle (Archiv der GPW oder Tagesstatistik der NewConnect): Kopfzeile mit „ISIN“ suchen, dann je ISIN
    Schlusskurs, Umsatz, Währung, Name."""
    head_at = next((k for k, row in enumerate(table) if any("isin" in str(c).lower() for c in row)), None)
    if head_at is None:
        return "ohne_isin_spalte", {}, []
    header = [" ".join(str(c or "").split()).lower() for c in table[head_at]]

    def col(*words: str) -> int | None:
        return next((k for k, c in enumerate(header) if any(w in c for w in words)), None)

    c_isin, c_close = col("isin"), col("zamkn", "closing", "close")
    c_cur, c_vol, c_name = col("walut", "currency"), col("wolumen", "volume"), col("nazwa", "name", "spółka", "instrument")
    if c_isin is None or c_close is None:
        return "spalten_fehlen", {}, [str(c) for c in table[head_at]]
    out: dict[str, dict] = {}
    for row in table[head_at + 1:]:
        cells = [("" if c is None else str(c)) for c in row]
        if len(cells) <= max(c_isin, c_close):
            continue
        isin = cells[c_isin].strip().upper()
        raw = row[c_close]
        close = float(raw) if isinstance(raw, (int, float)) else pl_number(cells[c_close])
        if not ISIN_RE.fullmatch(isin) or close is None or close <= 0:
            continue
        vol_raw = row[c_vol] if c_vol is not None and c_vol < len(row) else None
        out[isin] = {
            "close": close,
            "volume": float(vol_raw) if isinstance(vol_raw, (int, float)) else (pl_number(cells[c_vol]) if c_vol is not None and c_vol < len(cells) else None),
            "currency": cells[c_cur].strip() if c_cur is not None and c_cur < len(cells) else None,
            "name": cells[c_name].strip() if c_name is not None and c_name < len(cells) else None,
        }
    return ("ok" if out else "ohne_zeilen"), out, [str(c) for c in table[head_at]]


def sheet_tables(content: bytes, ext: str) -> list[list[list]]:
    """Alle Blätter einer Excel-Datei als Zeilen (Werte) – .xlsx mit openpyxl, .xls mit xlrd (im Job nachinstalliert)."""
    import io as _io

    if ext == "xlsx":
        import openpyxl

        book = openpyxl.load_workbook(_io.BytesIO(content), read_only=True, data_only=True)
        return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in book.worksheets]
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    return [[sh.row_values(k) for k in range(sh.nrows)] for sh in book.sheets()]


def fetch_newconnect(http, days: list[int], pause=time.sleep) -> tuple[Counter, dict[int, dict[str, dict]], list[str]]:
    """Tagesstatistik der NewConnect je Handelstag (.xlsx, sonst .xls); Abbruch wie bei der GPW nach drei Tagen ohne Daten."""
    status: Counter = Counter()
    by_day: dict[int, dict[str, dict]] = {}
    header: list[str] = []
    for n, day in enumerate(days):
        if n >= 3 and not by_day:
            status["abgebrochen_nach_3_tagen_ohne_daten"] += len(days) - n
            break
        ymd = (EPOCH + dt.timedelta(days=day)).strftime("%Y%m%d")
        result = "ohne_datei"
        for ext in ("xlsx", "xls"):
            code, content = welt.zusatz.get_raw(http, NC_URL.format(ymd=ymd, ext=ext), {"Accept": "*/*", "Referer": "https://newconnect.pl/archiwum-notowan"})
            if code != "ok":
                result = code
                continue
            try:
                tables = sheet_tables(content or b"", ext)
            except ImportError:
                result = f"modul_fehlt_{ext}"
                continue
            except Exception as exc:  # noqa: BLE001 - unlesbare Datei: nur zählen
                result = f"unlesbar_{ext}_{type(exc).__name__}"
                continue
            for table in tables:
                result, rows, head = parse_rows(table)
                header = header or head
                if result == "ok":
                    by_day[day] = rows
                    break
            if result == "ok":
                break
        status[result] += 1
    return status, by_day, header


def gpw_rows(by_day: dict[int, dict[str, dict]], isin: str) -> tuple[dict[int, dict], dict]:
    rows = {d: {"close": v[isin]["close"], "volume": v[isin]["volume"]} for d, v in by_day.items() if isin in v}
    sample = next((v[isin] for v in by_day.values() if isin in v), None)
    meta = {"isin": isin, "name": (sample or {}).get("name"), "boerse": VENUE_NAME["gpw"], "waehrung": (sample or {}).get("currency")} if sample else {}
    return rows, meta


def tmx_symbol(yahoo_symbol: str) -> str:
    """Yahoo-Kürzel → Kürzel bei TMX: Endung weg, Gattung mit Punkt („ABC-A.TO“ → „ABC.A“, „XYZ.V“ → „XYZ“)."""
    return welt.base_ticker(yahoo_symbol).replace("-", ".")


def tmx_post(http, query: str, variables: dict, name: str) -> tuple[str, object | None]:
    try:
        r = http.post(TMX_URL, json_body={"operationName": name, "variables": variables, "query": query}, headers=TMX_HEADERS)
    except Exception as exc:  # noqa: BLE001 - nur zählen
        return f"fehler_{type(exc).__name__}", None
    if r.status_code != 200:
        return f"http_{r.status_code}", None
    try:
        payload = r.json()
    except ValueError:
        return "unlesbar", None
    if not isinstance(payload, dict):
        return "unlesbar", None
    if payload.get("errors") and not (payload.get("data") or {}).get(name):
        return "graphql_fehler", payload.get("errors")
    return "ok", (payload.get("data") or {}).get(name)


def fetch_tmx(http, symbol: str, start_day: int, cutoff: int) -> tuple[str, dict, dict[int, dict], list[str]]:
    """Kursmeldung und Kurshistorie bei TMX: (Status, Angaben, Tageskurse, Fehlermeldungen der Schnittstelle ohne Werte)."""
    errors: list[str] = []
    code, quote = tmx_post(http, TMX_QUOTE, {"symbol": symbol, "locale": "en"}, "getQuoteBySymbol")
    if code == "graphql_fehler":
        errors += [str((e or {}).get("message", ""))[:120] for e in (quote or [])]
    if code != "ok" or not isinstance(quote, dict) or not quote.get("name"):
        return ("nicht_gefunden" if code == "ok" else f"kursmeldung_{code}"), {}, {}, errors
    meta = {"name": quote.get("name"), "boerse": quote.get("exchangeName") or quote.get("exchangeCode"),
            "waehrung": "USD" if symbol.upper().endswith(".U") else "CAD"}
    rows: dict[int, dict] = {}
    variables = {"symbol": symbol, "freq": "day", "interval": None, "start": iso(start_day), "end": iso(cutoff + 1), "startDateTime": None, "endDateTime": None}
    code, series = tmx_post(http, TMX_SERIES, variables, "getTimeSeriesData")
    if code == "graphql_fehler":
        errors += [str((e or {}).get("message", ""))[:120] for e in (series or [])]
    if code == "ok":
        rows = tmx_rows(series, "dateTime", "close", start_day, cutoff)
        meta["historie"] = "kursgrafik"
    if code != "ok" or len(rows) < 2:
        rows, end, pages = {}, cutoff, 0
        while end >= start_day and pages < TMX_PAGES:
            variables = {"symbol": symbol, "start": iso(start_day), "end": iso(end), "limit": 400, "adjusted": False, "unadjusted": True}
            code, history = tmx_post(http, TMX_HISTORY, variables, "getCompanyPriceHistory")
            pages += 1
            if code == "graphql_fehler":
                errors += [str((e or {}).get("message", ""))[:120] for e in (history or [])]
            if code != "ok":
                break
            part = tmx_rows(history, "datetime", "closePrice", start_day, cutoff)
            new = {d: r for d, r in part.items() if d not in rows}
            if not new:
                break
            rows.update(new)
            end = min(new) - 1
        meta["historie"] = f"kurshistorie_{pages}_abschnitte"
        if not rows:
            return f"historie_{code}", meta, {}, errors
    return ("ok" if rows else "keine_kurse"), meta, rows, errors


def tmx_rows(items, date_key: str, close_key: str, start_day: int, cutoff: int) -> dict[int, dict]:
    rows: dict[int, dict] = {}
    for item in items or []:
        if not isinstance(item, dict):
            continue
        text = str(item.get(date_key) or "")[:10]
        try:
            day = (dt.date.fromisoformat(text) - EPOCH).days
        except ValueError:
            continue
        close = item.get(close_key)
        if close is None or not (start_day <= day <= cutoff):
            continue
        rows[day] = {"close": float(close), "volume": item.get("volume")}
    return rows


def select_targets(stand, instruments: dict) -> tuple[list[int], dict[int, int], dict[int, int]]:
    """Aktien außerhalb der US-Auswahl mit Platz unter den ersten 10 an einer Börse mit eigener Quelle – die mit den meisten
    Plätzen 1 zuerst (Auswahl wie im Job „welt“)."""
    us = ursachen.us_stocks(instruments)
    chosen, top10, first = welt.select_world(stand, instruments, us)
    keys = instruments.get("exchange_key") or [None] * len(instruments["isin"])
    targets = [i for i in chosen if keys[i] in SOURCES]
    targets.sort(key=lambda i: (-first.get(i, 0), -top10.get(i, 0), i))
    return targets, top10, first


def run(stand, sid: str, http, figi_lookup=None, pause=time.sleep) -> dict:
    manifest = stand.get("manifest.bin")
    instruments = stand.get("instruments.bin")
    cutoff = int(manifest["cutoff_day"])
    min_trend = int((manifest.get("method") or {}).get("min_obs_trend", 15))
    liquid = (instruments.get("turnover") or {}).get("liquid") or [False] * len(instruments["isin"])
    start_day = cutoff - ursachen.gegen.LOOKBACK_DAYS - 5
    keys = instruments.get("exchange_key") or [None] * len(instruments["isin"])
    symbols = [str(s or "") for s in (instruments.get("symbol") or [None] * len(instruments["isin"]))]
    targets, top10, first = select_targets(stand, instruments)
    own = {i: ursachen.own_series(stand, instruments, i, start_day, cutoff) for i in targets}
    figi = (figi_lookup or absicherung.figi_lookup)([str(instruments["isin"][i]) for i in targets]) if targets else {}
    report: dict = {"stand": sid, "stichtag": nachrechnen._iso(cutoff), "aktien": [], "quellen": {}, "spalten_gpw": [], "tmx_fehler": []}
    war = [i for i in targets if keys[i] == "XWAR"]
    by_day: dict[int, dict[str, dict]] = {}
    if war:
        days = sorted({d for i in war for d in own[i]})
        status, by_day, header = fetch_gpw(http, days, pause=pause)
        report["quellen"]["gpw"] = dict(status)
        report["spalten_gpw"] = header
    tmx_status: Counter = Counter()
    nc_by_day: dict[int, dict[str, dict]] | None = None  # NewConnect erst laden, wenn eine Aktie nicht im Archiv des Hauptmarkts steht
    substituted = {}
    level_only: dict = {}  # nicht abgesichert, nur das Kursniveau weicht ab
    places = ursachen.best_places(stand, set(targets))
    for k, i in enumerate(targets):
        source = SOURCES[keys[i]]
        entry: dict = {"aktie": letter(k), "boerse": keys[i], "quelle": source, "top10": int(top10.get(i, 0)), "platz1": int(first.get(i, 0)),
                       "tage_stand": len(own[i])}
        found = places.get(i)
        entry["fundstelle"] = f"{found[1]}/{'AB'[found[2]]}/{found[3]} Platz {found[0]}" if found else "in keiner Rangliste"
        report["aktien"].append(entry)
        if not own[i]:
            entry["ergebnis"] = "stand_ohne_kurse"
            continue
        if source == "gpw":
            rows, meta_source = gpw_rows(by_day, str(instruments["isin"][i]))
            code = "ok" if rows else ("nicht_im_archiv" if by_day else "ohne_antwort")
            if code == "nicht_im_archiv":
                if nc_by_day is None:
                    nc_days = sorted({d for j in war for d in own[j]})
                    nc_status, nc_by_day, nc_header = fetch_newconnect(http, nc_days, pause=pause)
                    report["quellen"]["newconnect"] = dict(nc_status)
                    report["spalten_newconnect"] = nc_header
                rows, meta_source = gpw_rows(nc_by_day, str(instruments["isin"][i]))
                if rows:
                    meta_source["boerse"] = "Warsaw Stock Exchange (NewConnect)"
                    entry["quelle"] = "newconnect"
                    code = "ok"
                else:
                    code = "nicht_im_archiv_auch_nicht_newconnect"
        else:
            code, meta_source, rows, errors = fetch_tmx(http, tmx_symbol(symbols[i]), start_day, cutoff)
            tmx_status[code] += 1
            report["tmx_fehler"] += [e for e in errors if e not in report["tmx_fehler"]]
        if code != "ok":
            entry["ergebnis"] = code
            continue
        yahoo = ursachen.yahoo_daily(http, symbols[i], start_day, cutoff)
        checks, converted = welt.check_world(rows, meta_source, own[i], ursachen.stand_meta(instruments, i), yahoo, figi.get(str(instruments["isin"][i])),
                                             tmx_symbol(symbols[i]), symbols[i], keys[i], cutoff)
        checks["quelle"] = f"boerse_{source}"
        if source == "tmx" and checks.get("name") != "bestaetigt" and (absicherung.name_score(meta_source.get("name"), ursachen.stand_meta(instruments, i)["namen"]) or 0) >= absicherung.NAME_CONFIRMED:
            # TMX nennt die Gattung im Namen („… Class A Subordinate Voting Shares“) – verglichen wird wie bei ``name_score`` auch
            # der Name des Emittenten; die Gattung prüft OpenFIGI (Aktie und Gattung)
            checks["name"] = "bestaetigt"
            checks["abgesichert"], checks["grund"] = absicherung.verdict(checks)
        entry.update(ergebnis="abgesichert" if checks["abgesichert"] else "nicht_abgesichert", grund=checks["grund"],
                     kursniveau=checks.get("kursniveau_median"), gemeinsame_tage=checks.get("gemeinsame_tage"), historie=meta_source.get("historie"))
        if not checks["abgesichert"]:
            entry["befund"] = absicherung.failure_details(checks, figi.get(str(instruments["isin"][i])), keys[i])
            entry.update(factor_analysis(own[i], converted))
            if checks["grund"] == "kursniveau":
                # Nur das Kursniveau weicht ab (Aktie, Name, Börse, Währung, Handelstag, Bereinigung bestätigt): als eigene Variante
                # „Wenn die Börse recht hätte“ neu gerechnet – ausdrücklich nicht abgesichert, ersetzt nichts
                days = welt.classify_days(own[i], converted, yahoo, (yahoo or {}).get("splits") or [], cutoff, ursachen.within_display)
                deviating = {d: v for d, v in days.items() if v[0] not in ("gleich", "gleich_nach_split")}
                entry["abweichend_ungesichert"] = len(deviating)
                if deviating:
                    level_only[i] = ursachen.substitution(own[i], {d: v[1] for d, v in deviating.items()})
            continue
        days = welt.classify_days(own[i], converted, yahoo, (yahoo or {}).get("splits") or [], cutoff, ursachen.within_display)
        deviating = {d: v for d, v in days.items() if v[0] not in ("gleich", "gleich_nach_split")}
        entry["tage_verglichen"] = len(days)
        entry["gleich"] = sum(1 for v in days.values() if v[0] == "gleich")
        entry["split"] = sum(1 for v in days.values() if v[0] == "gleich_nach_split")
        entry["abweichend"] = len(deviating)
        entry["ursachen"] = dict(Counter(v[2] for v in deviating.values()))
        entry["groesste_abweichung_pct"] = max((round(100 * abs(own[i][d] / v[1] - 1), 3) for d, v in deviating.items()), default=0.0)
        entry["nur_im_stand"] = sum(1 for d in own[i] if d not in converted)
        if deviating:
            substituted[i] = ursachen.substitution(own[i], {d: v[1] for d, v in deviating.items()})
    if tmx_status:
        report["quellen"]["tmx"] = dict(tmx_status)
    variants = {"abgesichert": substituted, "mit_kursniveau": {**substituted, **level_only}}
    variants = {k: v for k, v in variants.items() if v and (k == "abgesichert" or level_only)}
    effects = ursachen.ranking_effects(stand, variants, liquid, min_trend, set(targets)) if variants else {}
    letter_of = {i: report["aktien"][k]["aktie"] for k, i in enumerate(targets)}

    def name(i) -> str:
        return letter_of.get(i, "eine nicht geprüfte Aktie") if i is not None else "–"

    def change_rows(changes: list[dict]) -> list[dict]:
        out = []
        for c in changes:
            if not c.get("platz1"):
                continue
            old, new = c["platz1"]
            out.append({"h": c["h"], "gruppe": c["gruppe"], "w": c["w"], "vorher": name(old), "nachher": name(new),
                        "ausloeser": None if c.get("ausloeser") is None else [name(j) for j in c["ausloeser"]],
                        "nicht_untersucht": c.get("nicht_untersucht") or False})
        return out

    report["wirkung"] = {}
    by_index = {i: report["aktien"][k] for k, i in enumerate(targets)}
    for variant, (totals, moves, changes) in effects.items():
        report["wirkung"][variant] = {"neu_gerechnet": len(variants[variant]), "zaehler": totals, "platz1_geaendert": change_rows(changes)}
        # ``moves`` zählt nur in Zeiträumen mit geänderten Werten – die übrigen Ranglisten bleiben, wie sie sind
        for i, m in moves.items():
            if i in by_index and i in variants[variant] and (variant == "abgesichert") == (i in substituted):
                a = by_index[i]
                a["platz1_nachher"] = a["platz1"] - int(m.get("platz1_vorher", 0)) + int(m.get("platz1_nachher", 0))
                a["top10_nachher"] = a["top10"] - int(m.get("unter_ersten_10_vorher", 0)) + int(m.get("unter_ersten_10_nachher", 0))
    return report


def letter(k: int) -> str:
    return ursachen.letter_label(k)


def simple_fraction(x: float, max_den: int = 20, tol: float = 0.001) -> str | None:
    """Einfacher Bruch p/q (q ≤ 20) nahe x – z. B. 18/17 oder 6/5 (Verhältnis einer Kapitalmaßnahme?); sonst None."""
    best = None
    for q in range(1, max_den + 1):
        p = round(x * q)
        if p > 0 and abs(p / q - x) <= tol * x and (best is None or q < best[1]):
            best = (p, q)
    return f"{best[0]}/{best[1]}" if best and best[1] > 1 else None


def factor_analysis(own: dict[int, float], rows: dict[int, dict]) -> dict:
    """Bei abweichendem Kursniveau: Ist das Verhältnis Stand/Quelle an allen gemeinsamen Tagen gleich (dann ändert es keine
    Rendite – Hinweis auf eine Bereinigung oder Einheit), oder ändert es sich ab einem Tag (dann ändern sich Renditen über diesen
    Tag)? Nur Kennzahlen, keine Kurse."""
    days = sorted(d for d in own if d in rows and rows[d]["close"])
    if len(days) < 2:
        return {}
    ratios = [own[d] / rows[d]["close"] for d in days]
    lo, hi = min(ratios), max(ratios)
    out = {"faktor_min": round(lo, 4), "faktor_max": round(hi, 4), "faktor_tage": len(days), "faktor_bruch": simple_fraction(sorted(ratios)[len(ratios) // 2])}
    out["faktor_konstant"] = (hi - lo) <= 0.005 * hi
    if not out["faktor_konstant"]:
        # erster Tag, ab dem das Verhältnis dauerhaft beim letzten Wert liegt (Wechsel der Bereinigung)
        last = ratios[-1]
        k = len(ratios) - 1
        while k > 0 and abs(ratios[k - 1] - last) <= 0.005 * last:
            k -= 1
        out["faktor_wechsel_tage_vor_stichtag"] = len(days) - k
    return out


RESULT_TEXT = {
    "abgesichert": "abgesichert verglichen", "nicht_abgesichert": "Kurse gefunden, Zuordnung nicht ausreichend abgesichert",
    "nicht_im_archiv": "ISIN nicht im Archiv der GPW (Hauptmarkt)", "ohne_antwort": "Quelle ohne Antwort", "stand_ohne_kurse": "Stand ohne Kurse im Zeitraum",
    "nicht_im_archiv_auch_nicht_newconnect": "ISIN weder im Archiv des Hauptmarkts noch in der Tagesstatistik der NewConnect",
    "nicht_gefunden": "bei TMX nicht gefunden", "keine_kurse": "TMX ohne Kurse im Zeitraum",
}


def result_text(code: str) -> str:
    if code in RESULT_TEXT:
        return RESULT_TEXT[code]
    if code.startswith("kursmeldung_") or code.startswith("historie_"):
        return f"TMX ohne verwertbare Antwort ({code})"
    return code


def lines(report: dict) -> tuple[str, list[str], list[str]]:
    rows = []
    for a in report["aktien"]:
        text = (f"{a['aktie']} = {a['fundstelle']} – {a['boerse']}, {SOURCE_TEXT[a['quelle']]}: Plätze unter den ersten 10 {a['top10']}, Platz 1 {a['platz1']}; "
                f"{result_text(a.get('ergebnis', '–'))}")
        if a.get("historie"):
            text += f" (Historie: {a['historie']})"
        if a.get("ergebnis") == "abgesichert":
            text += (f"; Tage im Stand {a['tage_stand']}, verglichen {a['tage_verglichen']}: gleich {a['gleich']}, nach Split gleich {a['split']}, "
                     f"abweichend {a['abweichend']}" + (f" (größte Abweichung {a['groesste_abweichung_pct']:.3f} %; Einordnung: "
                                                         + ", ".join(f"{ursachen._cause_label(c)} {n}" for c, n in sorted(a["ursachen"].items())) + ")" if a["abweichend"] else "")
                     + f"; nur im Stand {a['nur_im_stand']}")
            if a["abweichend"]:
                text += f"; neu gerechnet: Platz 1 {a.get('platz1_nachher', a['platz1'])} (vorher {a['platz1']}), unter den ersten 10 {a.get('top10_nachher', a['top10'])}"
        elif a.get("ergebnis") == "nicht_abgesichert":
            text += f" ({a.get('grund') or '–'}; Kursniveau Stand/Quelle im Median {a.get('kursniveau') if a.get('kursniveau') is not None else '–'}"
            text += f", gemeinsame Tage {a.get('gemeinsame_tage', 0)}) – {a.get('befund') or ''}".rstrip(" –")
            if a.get("abweichend_ungesichert") is not None:
                text += f"; Tage mit anderem Kurs als an der Börse {a['abweichend_ungesichert']}"
                if a.get("platz1_nachher") is not None:
                    text += f" (nur zur Einordnung mit den Kursen der Börse: Platz 1 {a['platz1_nachher']} statt {a['platz1']})"
            if a.get("faktor_tage"):
                text += (f"; Verhältnis Stand/Quelle an {a['faktor_tage']} gemeinsamen Tagen {a['faktor_min']} bis {a['faktor_max']}"
                         + (f" (≈ {a['faktor_bruch']})" if a.get("faktor_bruch") else "")
                         + (" – an allen Tagen gleich: Renditen im Zeitraum stimmen überein, nur das Kursniveau (Bereinigung oder Einheit) weicht ab"
                            if a.get("faktor_konstant") else f" – nicht gleich: wechselt {a.get('faktor_wechsel_tage_vor_stichtag')} Kurstage vor dem Stichtag (Renditen über diesen Tag weichen ab)"))
        rows.append(text)
    total = len(report["aktien"])
    secured = [a for a in report["aktien"] if a.get("ergebnis") == "abgesichert"]
    leaders = [a for a in report["aktien"] if a["platz1"]]
    leaders_ok = [a for a in leaders if a.get("ergebnis") == "abgesichert"]
    w = report.get("wirkung") or {}

    def effect_line(label: str, e: dict) -> str:
        changed = e.get("platz1_geaendert") or []
        rows_ = [f"{c['h']}/{c['gruppe']}/{c['w']}: {c['vorher']} → {c['nachher']}"
                 + (f" (Auslöser {', '.join(c['ausloeser'])})" if c.get("ausloeser") else (" (Auslöser nicht untersucht)" if c.get("nicht_untersucht") else ""))
                 for c in changed[:20]]
        return (f"{label}: {e.get('neu_gerechnet', 0)} Aktien neu gerechnet; Ranglisten mit anderem Platz 1 {len(changed)}"
                + (" – " + "; ".join(rows_) if rows_ else "")
                + f"; mit anderen Aktien unter den ersten 10 {int((e.get('zaehler') or {}).get('top10_aktien_geaendert', 0))}.")
    head = [
        f"Stand {report['stand']} (Kursstand {report['stichtag']}): {total} Aktien an XWAR, XTSE, XTSX mit Platz unter den ersten 10 "
        f"(davon {len(leaders)} mit Platz 1 in zusammen {sum(a['platz1'] for a in leaders)} Ranglisten) gegen die Quellen der Börsenbetreiber; "
        f"abgesichert verglichen {len(secured)} (mit Platz 1: {len(leaders_ok)}, zusammen {sum(a['platz1'] for a in leaders_ok)} Ranglisten).",
        "Abrufe: " + "; ".join(f"{SOURCE_TEXT[q]} " + ", ".join(f"{k} {v}" for k, v in sorted(s.items())) for q, s in sorted(report["quellen"].items())) if report["quellen"] else "Abrufe: –",
        effect_line("Wirkung (abgesicherte Vergleiche, Kurse der Börse an jedem abweichenden Tag)", w.get("abgesichert") or {}),
    ]
    if w.get("mit_kursniveau"):
        head.append(effect_line("Nur zur Einordnung, nicht abgesichert – zusätzlich die Aktien, bei denen nur das Kursniveau abweicht ("
                                + ", ".join(a["aktie"] for a in report["aktien"] if a.get("abweichend_ungesichert")) + "), mit den Kursen der Börse", w["mit_kursniveau"]))
    head += [
        "Geändert wird nichts: Die Vergleichskurse sind nachträglich beschafft und gelten nicht als Originaldaten; ein Unterschied zeigt "
        "nur, wo Stand und Börse auseinanderliegen.",
    ]
    if report.get("spalten_gpw"):
        head.append("Spalten des GPW-Archivs (nur Namen): " + ", ".join(report["spalten_gpw"]))
    if report.get("spalten_newconnect"):
        head.append("Spalten der NewConnect-Tagesstatistik (nur Namen): " + ", ".join(report["spalten_newconnect"]))
    if report.get("tmx_fehler"):
        head.append("Meldungen der TMX-Schnittstelle (ohne Werte): " + " | ".join(report["tmx_fehler"][:5]))
    title = f"Spitzenreiter an den Börsen geprüft: {len(secured)} von {total} abgesichert, Platz 1 in {sum(a['platz1'] for a in leaders_ok)} von {sum(a['platz1'] for a in leaders)} Ranglisten"
    return title, head, rows


def unresolved(report: dict) -> list[str]:
    out = []
    leaders = [a for a in report["aktien"] if a["platz1"] and a.get("ergebnis") != "abgesichert"]
    if leaders:
        out.append(f"{len(leaders)} Spitzenreiter ohne abgesicherten Vergleich ({', '.join(a['aktie'] for a in leaders)}; zusammen {sum(a['platz1'] for a in leaders)} Ranglisten mit Platz 1)")
    deviating = [a for a in report["aktien"] if a.get("abweichend")]
    if deviating:
        out.append(f"{len(deviating)} Aktien mit Abweichung zur Börse ({', '.join(a['aktie'] for a in deviating)})")
    changed = ((report.get("wirkung") or {}).get("abgesichert") or {}).get("platz1_geaendert") or []
    if changed:
        out.append(f"{len(changed)} Ranglisten hätten mit den Kursen der Börse einen anderen Platz 1")
    return out


def fx_probe(http, cutoff: int) -> list[str]:
    """Erreichbarkeit von Quellen für TWD und CLP vom Prüfrechner aus: Status, letzter Tag, Zahl der Tage in 30 Tagen vor dem
    Stichtag – keine Werte."""
    start = iso(cutoff - 45)
    probes = [
        ("Fed H.10 über FRED, TWD je USD", "https://fred.stlouisfed.org/graph/fredgraph.csv?id=DEXTAUS", "fred"),
        ("BIS, TWD je USD (täglich)", f"https://stats.bis.org/api/v1/data/WS_XRU_D/D.TW.TWD.A?startPeriod={start}&format=csv", "bis"),
        ("BIS, CLP je USD (täglich)", f"https://stats.bis.org/api/v1/data/WS_XRU_D/D.CL.CLP.A?startPeriod={start}&format=csv", "bis"),
        ("mindicador.cl (Banco Central de Chile), CLP je USD", "https://mindicador.cl/api/dolar", "mindicador"),
    ]
    out = []
    for label, url, kind in probes:
        code, content = welt.zusatz.get_raw(http, url, {"Accept": "text/csv, application/json, */*"})
        if code != "ok":
            out.append(f"{label}: {code}")
            continue
        days: list[int] = []
        text = (content or b"").decode("utf-8", errors="replace")
        try:
            if kind == "mindicador":
                import json

                for item in (json.loads(text) or {}).get("serie") or []:
                    days.append((dt.date.fromisoformat(str(item.get("fecha", ""))[:10]) - EPOCH).days)
            else:
                reader = csv.reader(io.StringIO(text))
                header = next(reader, [])
                col_d = next((k for k, h in enumerate(header) if h.lower() in ("observation_date", "date", "time_period")), 0)
                col_v = next((k for k, h in enumerate(header) if h.lower() in ("dextaus", "obs_value", "value")), len(header) - 1)
                for row in reader:
                    if len(row) > max(col_d, col_v) and pl_number(row[col_v]) not in (None, 0.0):
                        days.append((dt.date.fromisoformat(row[col_d][:10]) - EPOCH).days)
        except (ValueError, TypeError, AttributeError):
            out.append(f"{label}: Antwort nicht lesbar")
            continue
        recent = [d for d in days if cutoff - 30 <= d <= cutoff]
        out.append(f"{label}: erreichbar, letzter Tag {iso(max(days)) if days else '–'}, Tage in den 30 Tagen vor dem Stichtag {len(recent)}")
    return out


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--repo", default=os.environ.get("GITHUB_REPOSITORY", "szymeklh87-commits/tr-aktienvergleich"))
    parser.add_argument("--stand", default=os.environ.get("STAND") or None)
    args = parser.parse_args(argv)
    from traktien.providers.http import HttpClient

    import boersenplatz_auffaellige as venues

    started = time.monotonic()
    with tempfile.TemporaryDirectory() as temp:
        sid, folder, key = ursachen.load_stand(args.repo, os.environ.get("GITHUB_TOKEN"), Path(temp), args.stand)
        stand = nachrechnen.Stand(folder, sid, key)
        http = HttpClient(impersonate="chrome", max_retries=1, timeout=20, min_interval=HOSTS)
        report = run(stand, sid, http)
        fx = fx_probe(http, int(stand.get("manifest.bin")["cutoff_day"]))
    title, head, rows = lines(report)
    venues.notice(title, head + [f"Dauer {round(time.monotonic() - started)} s"])
    venues.notice("Spitzenreiter an den Börsen – je Aktie (Buchstaben statt Namen)", rows)
    venues.notice("Wechselkurse für TWD und CLP – frei zugängliche Quellen (Erreichbarkeit vom Prüfrechner)", fx)
    open_ = unresolved(report)
    if open_:
        print(f"::error title={ursachen.gegen._title('Spitzenreiter an den Börsen: nicht vollständig bestätigt')}::" + "%0A".join(open_))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
